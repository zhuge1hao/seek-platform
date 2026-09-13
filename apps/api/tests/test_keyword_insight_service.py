from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from services.keyword_insight_service import HEADERS, KeywordInsightError, analyze_rows, export_excel, infer_benchmarks, parse_percent, parse_range


def row(keyword: str, pop7: str, pop30: str | None, ctr7: str, ctr30: str | None, conv7: str, conv30: str | None, ctr_gap: float | None = None, conv_gap: float | None = None) -> dict:
    ctr_value = parse_percent(ctr7)
    conv_lower = parse_range(conv7).lower
    assert ctr_value is not None and conv_lower is not None
    return {
        "序号": 1,
        "搜索词": keyword,
        "近7天搜索人气": pop7,
        "近30天搜索人气": pop30,
        "近7天点击率": ctr7,
        "近30天点击率": ctr30,
        "点击率较平均值": abs(ctr_value - 102.4) if ctr_gap is None else ctr_gap,
        "近7天支付转化率": conv7,
        "近30天支付转化率": conv30,
        "转化率较平均值": abs(conv_lower - 22.5) if conv_gap is None else conv_gap,
    }


class KeywordInsightServiceTests(unittest.TestCase):
    def test_parsers_handle_units_percent_ranges_null_and_errors(self) -> None:
        self.assertEqual(parse_range("60万 ~ 120万").midpoint, 900000)
        self.assertEqual(parse_range("5000 ~ 1万").upper, 10000)
        self.assertEqual(parse_range(2500).lower, 2500)
        self.assertEqual(parse_range("5% ~ 7.5%").midpoint, 6.25)
        self.assertEqual(parse_percent("124%"), 124)
        self.assertIsNone(parse_range(None).midpoint)
        with self.assertRaises(KeywordInsightError):
            parse_range("异常")

    def test_benchmark_inference_uses_absolute_gap(self) -> None:
        rows = [
            row("甲", "5000 ~ 1万", "1万 ~ 2万", "124%", "120%", "30% ~ 35%", "30% ~ 35%"),
            row("乙", "2500 ~ 5000", "5000 ~ 1万", "44%", "50%", "5% ~ 7.5%", "5% ~ 7.5%"),
        ]
        inferred = infer_benchmarks(rows)
        self.assertEqual(inferred["ctr"]["value"], 102.4)
        self.assertEqual(inferred["conversion"]["value"], 22.5)
        self.assertEqual(inferred["ctr"]["confidence"], 1.0)

    def test_classifier_and_intent_cluster(self) -> None:
        context = {"industry": "女性内衣", "category": "女士内裤", "target_audience": ["妈妈"], "materials": ["纯棉"], "functions": ["高腰"]}
        rows = [
            row("女士高腰纯棉内裤", "2500 ~ 5000", "5000 ~ 1万", "119%", "124%", "30% ~ 35%", "30% ~ 35%"),
            row("三角裤", "4万 ~ 8万", "8万 ~ 15万", "23%", "32%", "5% ~ 7.5%", "7.5% ~ 10%"),
            row("孕妇内裤纯棉", "2500 ~ 5000", "1万 ~ 2万", "115%", "112%", "30% ~ 35%", "30% ~ 35%"),
            row("普通词", "1200 ~ 2500", None, "90%", None, "10% ~ 15%", None, ctr_gap=None, conv_gap=None),
        ]
        result = analyze_rows(rows, context)
        self.assertEqual([item["strategy"] for item in result["all_keywords"]], ["must_win", "blue_ocean_supply_gap", "blue_ocean_high_intent", "other"])
        variants = [row(name, "2500 ~ 5000", "1万 ~ 2万", "115%", "112%", "30% ~ 35%", "30% ~ 35%") for name in ("纯棉内裤女", "内裤女纯棉", "女士纯棉内裤")]
        clustered = analyze_rows(variants, context, {"deduplicate_intents": True})
        self.assertEqual(len({item["cluster_id"] for item in clustered["all_keywords"]}), 1)
        self.assertEqual(clustered["summary"]["high_intent_count"], 1)

    def test_export_preserves_four_sheets_headers_raw_values_and_order(self) -> None:
        rows = [row("先", "5000 ~ 1万", "1万 ~ 2万", "80%", "90%", "5% ~ 7.5%", "7.5% ~ 10%"), row("后", "2500 ~ 5000", None, "90%", None, "10% ~ 15%", None)]
        result = analyze_rows(rows)
        with tempfile.TemporaryDirectory() as directory:
            path = export_excel(result, Path(directory) / "out.xlsx")
            workbook = load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, ["全部搜索词", "行业必争搜索词", "供给不足蓝海词", "小众高意向蓝海词"])
            self.assertEqual([cell.value for cell in workbook["全部搜索词"][1]], HEADERS)
            self.assertEqual(workbook["全部搜索词"]["B2"].value, "先")
            self.assertEqual(workbook["全部搜索词"]["C2"].value, "5000 ~ 1万")
            self.assertIsNone(workbook["全部搜索词"]["D3"].value)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
