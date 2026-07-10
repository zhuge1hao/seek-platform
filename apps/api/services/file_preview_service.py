import csv
import hashlib
import json
import os
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from services import file_store
from services.artifact_service import build_download_url
from services.config_backup_service import resolve_runtime_path
from services.user_context import runtime_user_root


class FilePreviewError(RuntimeError):
    pass


def _safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _preview_excel(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        workbook.close()
        return {"columns": [], "rows": [], "row_count": 0, "preview_count": 0}

    columns = [_safe_text(value) or f"列{index + 1}" for index, value in enumerate(header)]
    rows: list[dict[str, str]] = []
    for row_index, row in enumerate(rows_iter):
        if row_index >= 20:
            break
        rows.append({columns[index]: _safe_text(value) for index, value in enumerate(row[: len(columns)])})

    row_count = max((sheet.max_row or 1) - 1, 0)
    workbook.close()
    return {"columns": columns, "rows": rows, "row_count": row_count, "preview_count": len(rows)}


def _preview_csv(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        columns = reader.fieldnames or []
        rows: list[dict[str, str]] = []
        row_count = 0
        for row in reader:
            row_count += 1
            if len(rows) < 20:
                rows.append({key: _safe_text(value) for key, value in row.items()})
    return {"columns": columns, "rows": rows, "row_count": row_count, "preview_count": len(rows)}


def _preview_text(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return {"text": text[:3000], "length": len(text)}


def _preview_json(path: Path) -> dict[str, Any]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    try:
        data = json.loads(raw)
        pretty = json.dumps(data, ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        data = None
        pretty = raw
    top_level = "array" if isinstance(data, list) else "object" if isinstance(data, dict) else type(data).__name__
    keys = list(data.keys())[:20] if isinstance(data, dict) else []
    return {"top_level": top_level, "keys": keys, "text": pretty[:3000], "length": len(pretty)}


def _cache_dir(user_id: str | None = None) -> Path:
    path = runtime_user_root(user_id) / "file_previews" if user_id else resolve_runtime_path(os.getenv("FILE_PREVIEW_CACHE_DIR", "runtime/file_previews"))
    path.mkdir(parents=True, exist_ok=True)
    return path


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cache_path(file_id: str, user_id: str | None = None) -> Path:
    return _cache_dir(user_id) / f"{file_id}.json"


def clear_preview_cache() -> int:
    count = 0
    cache_paths = list(_cache_dir().glob("*.json"))
    users_root = Path(__file__).resolve().parents[1] / "runtime" / "users"
    if users_root.exists():
        cache_paths.extend(users_root.glob("*/file_previews/*.json"))
    for path in cache_paths:
        path.unlink(missing_ok=True)
        count += 1
    return count


def preview_file(file_id: str, user_id: str | None = None, include_all_users: bool = False) -> dict[str, Any]:
    record = file_store.get_file_record(file_id, user_id=user_id, include_all_users=include_all_users, include_legacy=include_all_users)
    if record is None:
        raise FilePreviewError("文件记录不存在。")

    path = Path(record.get("saved_path") or "")
    if not path.exists():
        raise FilePreviewError("文件不存在或已被移动。")

    file_type = record.get("file_type")
    current_hash = _file_hash(path)
    cache_path = _cache_path(file_id, record.get("user_id") or user_id)
    if cache_path.exists():
        try:
            cached = json.loads(cache_path.read_text(encoding="utf-8"))
            if cached.get("sha256") == current_hash:
                if cached.get("error"):
                    raise FilePreviewError(cached["error"])
                return cached["response"]
        except json.JSONDecodeError:
            cache_path.unlink(missing_ok=True)

    try:
        if file_type == "excel":
            preview = _preview_csv(path) if path.suffix.lower() == ".csv" else _preview_excel(path)
        elif file_type == "text":
            preview = _preview_text(path)
        elif file_type == "json":
            preview = _preview_json(path)
        elif file_type in {"image", "video"}:
            preview = {"saved_path": str(path), "download_url": build_download_url(str(path))}
        else:
            preview = {"saved_path": str(path)}
    except Exception as exc:
        error = f"文件预览失败：{exc}"
        cache_path.write_text(json.dumps({"sha256": current_hash, "error": error}, ensure_ascii=False, indent=2), encoding="utf-8")
        raise FilePreviewError(error) from exc

    response = {
        "file_id": file_id,
        "file_type": file_type,
        "filename": record.get("filename"),
        "saved_path": record.get("saved_path"),
        "download_url": build_download_url(record.get("saved_path") or ""),
        "preview": preview,
    }
    cache_path.write_text(json.dumps({"sha256": current_hash, "response": response}, ensure_ascii=False, indent=2), encoding="utf-8")
    return response
