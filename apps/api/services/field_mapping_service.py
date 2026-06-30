import csv
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook

from services.user_context import datasets_dir


STANDARD_FIELDS: dict[str, str] = {
    "link": "链接", "index": "序号", "shop_name": "店铺名", "title": "产品标题",
    "gender": "性别", "age_group": "年龄", "consumption_level": "消费等级",
    "city_level": "城市等级", "product_category": "产品品类", "waist_type": "腰型结构",
    "main_material": "主面料", "material_function": "面料功能", "core_selling_point": "核心卖点",
    "secondary_selling_point": "辅助卖点", "design_element": "设计元素", "price": "价格",
    "total_price": "总价", "unit_price": "单条价", "spec_count": "规格条数",
    "sales_amount": "成交金额", "ad_spend": "付费花费", "cost_ratio": "费比",
    "sales_order_count": "成交订单数", "shipping_order_count": "发货订单数",
    "refund_rate": "退款率", "profit": "利润", "profit_rate": "利润率",
}

FIELD_ALIASES: dict[str, list[str]] = {
    "link": ["商品链接", "宝贝链接", "产品链接", "链接", "url", "商品url"],
    "index": ["序号", "编号", "排名", "index"],
    "shop_name": ["店铺", "店铺名", "店铺名称", "商家名称"],
    "title": ["产品标题", "商品标题", "宝贝标题", "标题", "商品名称"],
    "gender": ["性别", "适用性别"], "age_group": ["年龄", "年龄段", "适用年龄"],
    "consumption_level": ["消费等级", "消费层级"], "city_level": ["城市等级", "城市线级"],
    "product_category": ["产品品类", "商品类目", "品类", "类目"],
    "waist_type": ["腰型结构", "腰型"], "main_material": ["主面料", "材质", "面料"],
    "material_function": ["面料功能", "材质功能", "功能"],
    "core_selling_point": ["核心卖点", "主要卖点", "卖点"],
    "secondary_selling_point": ["辅助卖点", "次要卖点"], "design_element": ["设计元素", "设计点"],
    "price": ["价格", "商品价格"], "total_price": ["总价", "到手价", "售价", "成交价"],
    "unit_price": ["单条价", "单件价", "件单价", "单价"],
    "spec_count": ["规格条数", "条数", "件数", "规格数量"],
    "sales_amount": ["成交金额", "销售额", "成交额", "销售金额"],
    "ad_spend": ["付费花费", "推广费", "广告费", "站内推广费", "推广花费"],
    "cost_ratio": ["费比", "推广费占比", "站内推广费占比", "广告费占比"],
    "sales_order_count": ["成交订单数", "成交单量", "支付订单数"],
    "shipping_order_count": ["发货订单数", "发货单量"], "refund_rate": ["退款率", "退货率"],
    "profit": ["利润", "毛利"], "profit_rate": ["利润率", "毛利率"],
}


class TableReadError(RuntimeError):
    pass


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _unique_columns(values: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    columns: list[str] = []
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value is not None else ""
        base = base or f"列{index}"
        seen[base] = seen.get(base, 0) + 1
        columns.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return columns


def _rows_to_dicts(raw_rows: list[list[Any]]) -> tuple[list[str], list[dict[str, Any]]]:
    if not raw_rows:
        raise TableReadError("表格为空，无法创建数据集。")
    columns = _unique_columns(raw_rows[0])
    rows: list[dict[str, Any]] = []
    for values in raw_rows[1:]:
        padded = list(values) + [None] * max(0, len(columns) - len(values))
        if not any(value not in (None, "") for value in padded[:len(columns)]):
            continue
        rows.append({column: _json_value(padded[index]) for index, column in enumerate(columns)})
    return columns, rows


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            raw = [list(row) for row in sheet.iter_rows(values_only=True)]
            if raw and any(value not in (None, "") for value in raw[0]):
                return _rows_to_dicts(raw)
    finally:
        workbook.close()
    raise TableReadError("Excel 中没有可读取的工作表。")


def _read_xls(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = xlrd.open_workbook(path)
    for sheet in workbook.sheets():
        raw = [sheet.row_values(index) for index in range(sheet.nrows)]
        if raw and any(value not in (None, "") for value in raw[0]):
            return _rows_to_dicts(raw)
    raise TableReadError("Excel 中没有可读取的工作表。")


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as source:
                return _rows_to_dicts([list(row) for row in csv.reader(source)])
        except UnicodeDecodeError as exc:
            last_error = exc
    raise TableReadError(f"CSV 编码无法识别：{last_error}")


def read_table(path: str | Path) -> tuple[list[str], list[dict[str, Any]]]:
    source = Path(path)
    if not source.exists():
        raise TableReadError("源文件不存在。")
    try:
        if source.suffix.lower() == ".xlsx":
            return _read_xlsx(source)
        if source.suffix.lower() == ".xls":
            return _read_xls(source)
        if source.suffix.lower() == ".csv":
            return _read_csv(source)
    except TableReadError:
        raise
    except Exception as exc:
        raise TableReadError(f"表格读取失败：{exc}") from exc
    raise TableReadError("仅支持 xlsx、xls、csv 数据文件。")


def _normalize(value: str) -> str:
    return re.sub(r"[\s_\-（）()【】\[\]/\\]+", "", value.strip().lower())


def suggest_mapping(columns: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    used: set[str] = set()
    normalized_columns = {column: _normalize(column) for column in columns}
    for field, label in STANDARD_FIELDS.items():
        aliases = [_normalize(label), *(_normalize(alias) for alias in FIELD_ALIASES.get(field, []))]
        exact = next((column for column, normalized in normalized_columns.items() if column not in used and normalized in aliases), None)
        if exact:
            mapping[field] = exact
            used.add(exact)
            continue
        candidates = [column for column, normalized in normalized_columns.items() if column not in used and len(normalized) >= 2 and any(len(alias) >= 2 and (alias in normalized or normalized in alias) for alias in aliases)]
        if len(candidates) == 1:
            mapping[field] = candidates[0]
            used.add(candidates[0])
    return mapping


def validate_mapping(mapping: dict[str, str], columns: list[str]) -> dict[str, str]:
    unknown_fields = [field for field in mapping if field not in STANDARD_FIELDS]
    if unknown_fields:
        raise ValueError(f"包含未知标准字段：{'、'.join(unknown_fields)}")
    missing_columns = [column for column in mapping.values() if column and column not in columns]
    if missing_columns:
        raise ValueError(f"原始字段不存在：{'、'.join(missing_columns)}")
    selected = [column for column in mapping.values() if column]
    if len(selected) != len(set(selected)):
        raise ValueError("同一个原始字段不能映射到多个标准字段。")
    return {field: column for field, column in mapping.items() if column}


def preview_payload(path: str | Path, limit: int = 20) -> dict[str, Any]:
    columns, rows = read_table(path)
    return {
        "columns": columns,
        "rows": rows[:limit],
        "row_count": len(rows),
        "preview_count": min(limit, len(rows)),
        "detected_fields": [{"key": key, "name": name} for key, name in STANDARD_FIELDS.items()],
        "suggested_mapping": suggest_mapping(columns),
    }


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def list_mapping_templates(user_id: str) -> list[dict[str, Any]]:
    path = datasets_dir(user_id) / "mapping_templates.json"
    if not path.exists():
        return []
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
        items = value.get("templates", []) if isinstance(value, dict) else value
        return items if isinstance(items, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def save_mapping_template(user_id: str, name: str, mapping: dict[str, str]) -> dict[str, Any]:
    cleaned_name = name.strip()
    if not cleaned_name:
        raise ValueError("模板名称不能为空。")
    now = datetime.now().isoformat()
    template_id = f"mapping_{abs(hash(cleaned_name)) & 0xFFFFFFFF:08x}"
    item = {"template_id": template_id, "name": cleaned_name, "mapping": mapping, "updated_at": now}
    items = [existing for existing in list_mapping_templates(user_id) if existing.get("name") != cleaned_name]
    items.append(item)
    write_json(datasets_dir(user_id) / "mapping_templates.json", {"schema_version": "1.0", "updated_at": now, "templates": items})
    return item
