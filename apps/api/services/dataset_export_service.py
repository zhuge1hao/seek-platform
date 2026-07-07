import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from services.artifact_service import normalize_artifact_file
from services.field_mapping_service import write_json


def _cell_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6D5DD3")
    for index, column in enumerate(sheet.columns, start=1):
        width = max((len(str(cell.value or "")) for cell in column), default=8) + 2
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width, 10), 40)


def _write_rows(sheet, rows: list[dict[str, Any]]) -> None:
    columns: list[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    if not columns:
        columns = ["暂无数据"]
    sheet.append(columns)
    for row in rows:
        sheet.append([_cell_value(row.get(column)) for column in columns])
    _style_sheet(sheet)


def _overview_rows(profile: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        {"指标": "原始条数", "数值": profile.get("raw_count")},
        {"指标": "有效条数", "数值": profile.get("valid_count")},
        {"指标": "剔除条数", "数值": profile.get("excluded_count")},
        {"指标": "有效率", "数值": profile.get("valid_rate")},
        {"指标": "成交总额", "数值": profile.get("total_sales_amount")},
        {"指标": "平均费比", "数值": profile.get("avg_cost_ratio")},
        {"指标": "平均单条价", "数值": profile.get("avg_unit_price")},
    ]


def export_results(cleaned_dir: Path, valid_rows: list[dict[str, Any]], excluded_rows: list[dict[str, Any]], profile: dict[str, Any], metrics_summary: dict[str, Any]) -> list[dict[str, str]]:
    cleaned_dir.mkdir(parents=True, exist_ok=True)
    cleaned_path = cleaned_dir / "cleaned_data.xlsx"
    excluded_path = cleaned_dir / "excluded_data.xlsx"

    workbook = Workbook()
    valid_sheet = workbook.active
    valid_sheet.title = "有效数据"
    _write_rows(valid_sheet, valid_rows)
    overview = workbook.create_sheet("数据概览")
    _write_rows(overview, _overview_rows(profile))
    workbook.save(cleaned_path)

    excluded_book = Workbook()
    excluded_sheet = excluded_book.active
    excluded_sheet.title = "剔除数据"
    _write_rows(excluded_sheet, excluded_rows)
    reasons = excluded_book.create_sheet("剔除原因统计")
    _write_rows(reasons, [{"剔除原因": key, "数量": value} for key, value in profile.get("excluded_reason_counts", {}).items()])
    excluded_book.save(excluded_path)

    profile_path = cleaned_dir / "data_profile.json"
    metrics_path = cleaned_dir / "metrics_summary.json"
    write_json(profile_path, profile)
    write_json(metrics_path, metrics_summary)
    return [normalize_artifact_file(str(path)) for path in (cleaned_path, excluded_path, profile_path, metrics_path)]
