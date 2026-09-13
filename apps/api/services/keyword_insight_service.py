from __future__ import annotations

import hashlib
import json
import math
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


HEADERS = [
    "序号",
    "搜索词",
    "近7天搜索人气",
    "近30天搜索人气",
    "近7天点击率",
    "近30天点击率",
    "点击率较平均值",
    "近7天支付转化率",
    "近30天支付转化率",
    "转化率较平均值",
]
REQUIRED_HEADERS = [header for header in HEADERS if header != "序号"]
SHEETS = {
    "全部搜索词": None,
    "行业必争搜索词": "must_win",
    "供给不足蓝海词": "blue_ocean_supply_gap",
    "小众高意向蓝海词": "blue_ocean_high_intent",
}
RULES_PATH = Path(__file__).with_name("keyword_insight_rules.json")


class KeywordInsightError(ValueError):
    pass


@dataclass(frozen=True)
class RangeValue:
    raw: Any
    lower: float | None
    upper: float | None
    midpoint: float | None


def load_rules(path: str | Path | None = None) -> dict[str, Any]:
    return json.loads(Path(path or RULES_PATH).read_text(encoding="utf-8"))


def _number(token: str) -> float:
    token = token.strip().lower()
    multiplier = 10_000 if token.endswith("万") else 1
    if multiplier != 1:
        token = token[:-1]
    return float(token) * multiplier


def parse_range(value: Any) -> RangeValue:
    if value is None or str(value).strip() == "":
        return RangeValue(value, None, None, None)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        return RangeValue(value, number, number, number)
    text = str(value).strip().replace("％", "%")
    tokens = re.findall(r"\d+(?:\.\d+)?(?:万)?", text)
    if not tokens:
        raise KeywordInsightError(f"无法解析区间值：{value}")
    numbers = [_number(token) for token in tokens]
    lower, upper = numbers[0], numbers[-1]
    if lower > upper:
        raise KeywordInsightError(f"区间下界大于上界：{value}")
    return RangeValue(value, lower, upper, (lower + upper) / 2)


def parse_percent(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip().replace("％", "%").removesuffix("%")
    try:
        return float(text)
    except ValueError as exc:
        raise KeywordInsightError(f"无法解析百分比：{value}") from exc


def _optional_number(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise KeywordInsightError(f"无法解析数值：{value}") from exc


def read_excel_rows(source_file: str | Path, sheet_name: str = "全部搜索词") -> list[dict[str, Any]]:
    path = Path(source_file)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise KeywordInsightError("仅支持 .xlsx 或 .xlsm 文件。")
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise KeywordInsightError(f"缺少工作表：{sheet_name}")
    try:
        values = workbook[sheet_name].values
        headers = [str(value).strip() if value is not None else "" for value in next(values)]
        missing = [header for header in REQUIRED_HEADERS if header not in headers]
        if missing:
            raise KeywordInsightError(f"缺少关键字段：{', '.join(missing)}")
        rows = []
        for index, values_row in enumerate(values, 1):
            row = dict(zip(headers, values_row))
            if not any(value not in (None, "") for value in row.values()):
                continue
            row.setdefault("序号", index)
            if row.get("序号") in (None, ""):
                row["序号"] = index
            rows.append(row)
        return rows
    finally:
        workbook.close()


def read_ground_truth(source_file: str | Path) -> tuple[list[dict[str, Any]], dict[str, str], dict[str, int], list[str]]:
    workbook = load_workbook(source_file, read_only=True, data_only=True)
    sheet_names = workbook.sheetnames
    workbook.close()
    rows = read_excel_rows(source_file)
    labels: dict[str, str] = {}
    counts: dict[str, int] = {}
    overlaps: list[str] = []
    for sheet_name, label in SHEETS.items():
        if sheet_name not in sheet_names:
            raise KeywordInsightError(f"缺少工作表：{sheet_name}")
        sheet_rows = read_excel_rows(source_file, sheet_name)
        counts[sheet_name] = len(sheet_rows)
        if label is None:
            continue
        for row in sheet_rows:
            keyword = str(row["搜索词"]).strip()
            if keyword in labels and keyword not in overlaps:
                overlaps.append(keyword)
            labels[keyword] = label
    return rows, labels, counts, overlaps


def infer_benchmark(rows: list[dict[str, Any]], value_key: str, gap_key: str, *, range_lower: bool = False, tolerance: float = 0.11) -> dict[str, Any]:
    pairs: list[tuple[float, float]] = []
    candidates: Counter[float] = Counter()
    for row in rows:
        value = parse_range(row.get(value_key)).lower if range_lower else parse_percent(row.get(value_key))
        gap = _optional_number(row.get(gap_key))
        if value is None or gap is None:
            continue
        pairs.append((value, gap))
        candidates[round(value - gap, 4)] += 1
        candidates[round(value + gap, 4)] += 1
    if not pairs:
        return {"value": None, "confidence": 0.0, "supported_rows": 0, "valid_rows": 0, "max_error": None, "status": "unavailable"}
    benchmark = candidates.most_common(1)[0][0]
    errors = [abs(abs(value - benchmark) - gap) for value, gap in pairs]
    supported = sum(error <= tolerance for error in errors)
    confidence = supported / len(pairs)
    return {
        "value": benchmark,
        "confidence": round(confidence, 6),
        "supported_rows": supported,
        "valid_rows": len(pairs),
        "max_error": round(max(errors), 6),
        "status": "inferred" if confidence >= 0.95 else "low_confidence",
    }


def infer_benchmarks(rows: list[dict[str, Any]], rules: dict[str, Any] | None = None) -> dict[str, Any]:
    config = rules or load_rules()
    tolerance = float(config["benchmark_tolerance"])
    return {
        "ctr": infer_benchmark(rows, "近7天点击率", "点击率较平均值", tolerance=tolerance),
        "conversion": infer_benchmark(rows, "近7天支付转化率", "转化率较平均值", range_lower=True, tolerance=tolerance),
    }


def _daily_momentum(current: RangeValue, history: RangeValue) -> float | None:
    current_midpoint, history_midpoint = current.midpoint, history.midpoint
    if current_midpoint is None or history_midpoint in (None, 0):
        return None
    assert history_midpoint is not None
    return (current_midpoint / 7) / (history_midpoint / 30)


def _ratio(numerator: float | None, denominator: float | None) -> float | None:
    if numerator is None or denominator in (None, 0):
        return None
    assert denominator is not None
    return numerator / denominator


def parse_intent(keyword: str, rules: dict[str, Any]) -> dict[str, list[str]]:
    normalized = keyword.lower().replace(" ", "")
    return {
        group: [term for term in terms if term.lower() in normalized]
        for group, terms in rules["intent_lexicon"].items()
        if any(term.lower() in normalized for term in terms)
    }


def business_relevance(keyword: str, tags: dict[str, list[str]], context: dict[str, Any]) -> float:
    if not context:
        return 0.0
    normalized = keyword.lower().replace(" ", "")
    fields = {
        "category": [context.get("category"), context.get("product"), context.get("industry")],
        "audience": context.get("target_audience") or [],
        "material": context.get("materials") or [],
        "function": context.get("functions") or [],
        "style": context.get("styles") or [],
        "selling": context.get("core_selling_points") or [],
    }
    weights = {"category": 0.4, "audience": 0.15, "material": 0.15, "function": 0.15, "style": 0.05, "selling": 0.1}
    score = 0.0
    for field, values in fields.items():
        if isinstance(values, str) or values is None:
            values = [values]
        tokens = [str(value).lower().replace(" ", "") for value in values if value]
        if any(token in normalized or normalized in token for token in tokens):
            score += weights[field]
    if tags.get("category") and score < weights["category"]:
        score += weights["category"]
    return round(min(score, 1.0), 4)


def _canonical_keyword(keyword: str, tags: dict[str, list[str]]) -> str:
    groups = ["category", "audience", "material", "function", "style", "season", "scene", "size", "brand"]
    aliases = {"女士": "女", "女性": "女", "女生": "女", "女式": "女", "女款": "女", "全棉": "纯棉", "棉": "纯棉"}
    tokens = sorted({aliases.get(token.lower(), token.lower()) for group in groups for token in tags.get(group, [])})
    return "|".join(tokens) or re.sub(r"\W+", "", keyword.lower())


def _feature_row(row: dict[str, Any], ctr_benchmark: float, conversion_benchmark: float, rules: dict[str, Any], context: dict[str, Any]) -> dict[str, Any]:
    keyword = str(row.get("搜索词") or "").strip()
    if not keyword:
        raise KeywordInsightError("搜索词不能为空。")
    pop7 = parse_range(row.get("近7天搜索人气"))
    pop30 = parse_range(row.get("近30天搜索人气"))
    conv7 = parse_range(row.get("近7天支付转化率"))
    conv30 = parse_range(row.get("近30天支付转化率"))
    ctr7 = parse_percent(row.get("近7天点击率"))
    ctr30 = parse_percent(row.get("近30天点击率"))
    tags = parse_intent(keyword, rules)
    relevance = business_relevance(keyword, tags, context)
    return {
        **row,
        "keyword": keyword,
        "popularity_7d_lower": pop7.lower,
        "popularity_7d_upper": pop7.upper,
        "popularity_7d_mid": pop7.midpoint,
        "popularity_30d_lower": pop30.lower,
        "popularity_30d_upper": pop30.upper,
        "popularity_30d_mid": pop30.midpoint,
        "ctr_7d": ctr7,
        "ctr_30d": ctr30,
        "conversion_7d_lower": conv7.lower,
        "conversion_7d_upper": conv7.upper,
        "conversion_7d_mid": conv7.midpoint,
        "conversion_30d_lower": conv30.lower,
        "conversion_30d_upper": conv30.upper,
        "conversion_30d_mid": conv30.midpoint,
        "ctr_vs_avg_signed": None if ctr7 is None else ctr7 - ctr_benchmark,
        "conversion_vs_avg_signed": None if conv7.lower is None else conv7.lower - conversion_benchmark,
        "ctr_gap_abs": _optional_number(row.get("点击率较平均值")),
        "conversion_gap_abs": _optional_number(row.get("转化率较平均值")),
        "demand_momentum": _daily_momentum(pop7, pop30),
        "ctr_momentum": _ratio(ctr7, ctr30),
        "conversion_momentum": _ratio(conv7.midpoint, conv30.midpoint),
        "intent_tags": tags,
        "business_relevance_score": relevance,
        "cluster_id": hashlib.sha1(_canonical_keyword(keyword, tags).encode("utf-8")).hexdigest()[:12],
    }


def _score(feature: dict[str, Any], ctr_benchmark: float, conversion_benchmark: float) -> dict[str, float]:
    demand = min(100.0, 20.0 * math.log10(max(float(feature["popularity_7d_mid"] or 1), 1)))
    click = min(100.0, max(0.0, 50.0 + 2.0 * float((feature["ctr_7d"] or ctr_benchmark) - ctr_benchmark)))
    conversion = min(100.0, max(0.0, 50.0 + 4.0 * float((feature["conversion_7d_lower"] or conversion_benchmark) - conversion_benchmark)))
    growth = min(100.0, max(0.0, 50.0 * float(feature["demand_momentum"] or 1)))
    relevance = 100.0 * float(feature["business_relevance_score"])
    specificity = min(100.0, 20.0 * sum(bool(values) for values in feature["intent_tags"].values()))
    supply = 0.45 * demand + 0.3 * (100 - click) + 0.25 * (100 - conversion)
    high = 0.2 * demand + 0.3 * click + 0.35 * conversion + 0.15 * specificity
    must = 0.25 * demand + 0.2 * click + 0.2 * conversion + 0.2 * relevance + 0.15 * growth
    return {key: round(value, 2) for key, value in {"demand": demand, "click_health": click, "conversion_health": conversion, "growth": growth, "business_relevance": relevance, "intent_specificity": specificity, "supply_gap": supply, "high_intent": high, "must_win": must}.items()}


def _classify(feature: dict[str, Any], rules: dict[str, Any]) -> tuple[str, list[str], float]:
    must = rules["must_win"]
    supply = rules["supply_gap"]
    high = rules["high_intent"]
    momentum = feature["demand_momentum"]
    ctr7, ctr30 = feature["ctr_7d"], feature["ctr_30d"]
    conv7lo, conv7mid, conv30lo = feature["conversion_7d_lower"], feature["conversion_7d_mid"], feature["conversion_30d_lower"]
    pop7lo, pop7mid = feature["popularity_7d_lower"], feature["popularity_7d_mid"]

    if (
        momentum is not None
        and momentum > must["demand_momentum_min"]
        and ctr30 is not None
        and ctr30 > must["ctr_30d_min"]
        and conv30lo is not None
        and conv30lo > must["conversion_30d_lower_min"]
        and feature["business_relevance_score"] >= must["business_relevance_min"]
    ):
        return "must_win", ["近7天日均需求相对30天提升", "点击与支付转化健康", "匹配业务上下文"], 0.88

    weak_supply = pop7lo is not None and pop7lo >= supply["min_popularity_7d_lower"] and (
        (ctr7 is not None and ctr7 <= supply["weak_ctr_max"])
        or (conv7lo is not None and conv7lo <= supply["weak_conversion_lower_max"] and ctr30 is not None and ctr30 > supply["supporting_ctr_30d_min"])
    )
    if weak_supply:
        reasons = ["搜索需求达到供给不足候选规模"]
        if ctr7 is not None and ctr7 <= supply["weak_ctr_max"]:
            reasons.append("点击率明显低于行业参考")
        if conv7lo is not None and conv7lo <= supply["weak_conversion_lower_max"]:
            reasons.append("支付转化率低于健康区间")
        return "blue_ocean_supply_gap", reasons, 0.84

    healthy_niche = (
        conv30lo is not None
        and conv30lo > high["conversion_30d_lower_min"]
        and pop7mid is not None
        and pop7mid <= high["max_popularity_7d_mid"]
        and ctr7 is not None
        and ctr7 > high["ctr_7d_min"]
    )
    recovered_niche = (
        conv30lo is not None
        and conv30lo <= high["conversion_30d_lower_min"]
        and conv7mid is not None
        and conv7mid > high["recovery_conversion_7d_mid_min"]
        and ctr30 is not None
        and ctr30 > high["recovery_ctr_30d_min"]
    )
    if healthy_niche or recovered_niche:
        return "blue_ocean_high_intent", ["搜索规模处于细分区间", "点击率高于行业参考", "支付转化表现健康"], 0.86
    return "other", ["未同时满足三类策略规则"], 0.65


def _deduplicate(features: list[dict[str, Any]]) -> list[dict[str, Any]]:
    keep: dict[tuple[str, str], dict[str, Any]] = {}
    for feature in features:
        strategy = feature["strategy"]
        if strategy == "other":
            continue
        key = (strategy, feature["cluster_id"])
        rank = (feature["score"], feature["popularity_7d_mid"] or 0, len(feature["keyword"]))
        current = keep.get(key)
        if current is None or rank > (current["score"], current["popularity_7d_mid"] or 0, len(current["keyword"])):
            keep[key] = feature
    selected = {id(feature) for feature in keep.values()}
    for feature in features:
        if feature["strategy"] != "other" and id(feature) not in selected:
            feature["strategy"] = "other"
            feature["reason"] = ["同一意图簇已有更高分代表词"]
    return features


def analyze_rows(rows: Iterable[dict[str, Any]], business_context: dict[str, Any] | None = None, options: dict[str, Any] | None = None, rules: dict[str, Any] | None = None) -> dict[str, Any]:
    source_rows = [dict(row) for row in rows]
    missing = [header for header in REQUIRED_HEADERS if source_rows and header not in source_rows[0]]
    if missing:
        raise KeywordInsightError(f"缺少关键字段：{', '.join(missing)}")
    config = rules or load_rules()
    benchmarks = infer_benchmarks(source_rows, config)
    if benchmarks["ctr"]["value"] is None or benchmarks["conversion"]["value"] is None:
        raise KeywordInsightError("无法从有效的平均值差距字段推断行业基准。")
    ctr_benchmark = float(benchmarks["ctr"]["value"])
    conversion_benchmark = float(benchmarks["conversion"]["value"])
    features = []
    for index, row in enumerate(source_rows, 1):
        row.setdefault("序号", index)
        feature = _feature_row(row, ctr_benchmark, conversion_benchmark, config, business_context or {})
        strategy, reason, confidence = _classify(feature, config)
        scores = _score(feature, ctr_benchmark, conversion_benchmark)
        feature.update(strategy=strategy, score=scores[{"must_win": "must_win", "blue_ocean_supply_gap": "supply_gap", "blue_ocean_high_intent": "high_intent", "other": "demand"}[strategy]], scores=scores, reason=reason, confidence=confidence)
        features.append(feature)
    if (options or {}).get("deduplicate_intents", False):
        _deduplicate(features)
    limits = {"must_win": "max_must_win", "blue_ocean_supply_gap": "max_supply_gap", "blue_ocean_high_intent": "max_high_intent"}
    for strategy, option_name in limits.items():
        limit = (options or {}).get(option_name)
        if limit is not None:
            ranked = sorted((feature for feature in features if feature["strategy"] == strategy), key=lambda item: (-item["score"], item["序号"]))
            for feature in ranked[int(limit):]:
                feature["strategy"] = "other"
                feature["reason"] = [f"超过配置的 {option_name} 上限"]
    groups = {strategy: [feature for feature in features if feature["strategy"] == strategy] for strategy in ("must_win", "blue_ocean_supply_gap", "blue_ocean_high_intent")}
    return {
        "summary": {"total_keywords": len(features), "must_win_count": len(groups["must_win"]), "supply_gap_count": len(groups["blue_ocean_supply_gap"]), "high_intent_count": len(groups["blue_ocean_high_intent"])},
        "benchmarks": benchmarks,
        "must_win_keywords": groups["must_win"],
        "supply_gap_keywords": groups["blue_ocean_supply_gap"],
        "high_intent_keywords": groups["blue_ocean_high_intent"],
        "all_keywords": features,
    }


def analyze_excel(source_file: str | Path, business_context: dict[str, Any] | None = None, options: dict[str, Any] | None = None) -> dict[str, Any]:
    return analyze_rows(read_excel_rows(source_file), business_context, options)


def export_excel(result: dict[str, Any], output_file: str | Path, include_explanation_sheet: bool = False) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    groups = {
        "全部搜索词": result["all_keywords"],
        "行业必争搜索词": result["must_win_keywords"],
        "供给不足蓝海词": result["supply_gap_keywords"],
        "小众高意向蓝海词": result["high_intent_keywords"],
    }
    for sheet_name, rows in groups.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(HEADERS)
        for index, row in enumerate(rows, 1):
            sheet.append([index if sheet_name != "全部搜索词" else row.get("序号", index), *[row.get(header) for header in HEADERS[1:]]])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:J{max(1, sheet.max_row)}"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(right=Side(style="thin", color="FFFFFF"))
        for row_cells in sheet.iter_rows(min_row=2, min_col=3, max_col=10):
            for cell in row_cells:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        widths = [8, 34, 20, 20, 16, 16, 20, 22, 22, 20]
        for column, width in zip("ABCDEFGHIJ", widths):
            sheet.column_dimensions[column].width = width
    if include_explanation_sheet:
        sheet = workbook.create_sheet("关键词洞察说明")
        sheet.append(["搜索词", "策略", "得分", "置信度", "原因", "意图簇"])
        for row in result["all_keywords"]:
            sheet.append([row["keyword"], row["strategy"], row["score"], row["confidence"], "；".join(row["reason"]), row["cluster_id"]])
    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def classification_metrics(result: dict[str, Any], truth: dict[str, str]) -> dict[str, Any]:
    mapping = {"must_win": "must_win", "blue_ocean_supply_gap": "blue_ocean_supply_gap", "blue_ocean_high_intent": "blue_ocean_high_intent"}
    predictions = {row["keyword"]: row["strategy"] for row in result["all_keywords"]}
    metrics: dict[str, Any] = {}
    for label in mapping:
        truth_set = {keyword for keyword, value in truth.items() if value == label}
        predicted_set = {keyword for keyword, value in predictions.items() if value == label}
        tp, fp, fn = truth_set & predicted_set, predicted_set - truth_set, truth_set - predicted_set
        precision = len(tp) / len(predicted_set) if predicted_set else 0.0
        recall = len(tp) / len(truth_set) if truth_set else 0.0
        f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
        metrics[label] = {"true_positive": len(tp), "false_positive": len(fp), "false_negative": len(fn), "precision": round(precision, 6), "recall": round(recall, 6), "f1": round(f1, 6), "false_positive_keywords": sorted(fp), "false_negative_keywords": sorted(fn)}
    return metrics
